#!/usr/bin/env python3
"""Navigation plan tooling: workbook -> validated plan -> diff against live SharePoint navigation.

Purpose: the offline (no tenant access) half of the sp-navigation-sync skill. It validates
inventory/navigation-menu.xlsx, turns it into a plan, compares the plan with a live snapshot
produced by export-spo-navigation.ps1, and writes the operations list that
sync-spo-navigation.ps1 executes. It can also turn a snapshot into an editable xlsx/csv.

Key Input Dependencies:
- openpyxl (xlsx read/write).
- A workbook with headers "Top Level", "2nd Level", "3rd Level" and optional "Url",
  "OpenInNewTab", "Enabled" (first sheet is read).
- A snapshot JSON: {"siteUrl", "location", "topNav"|"quickLaunch": [{id,title,url,children}]}.

Usage:
    python navigation_plan.py plan --workbook nav.xlsx --out plan.json [--allow-missing-urls]
    python navigation_plan.py diff --plan plan.json --snapshot live.json --out ops.json [--site-url URL]
    python navigation_plan.py to-workbook --snapshot live.json --out nav.xlsx   (or .csv)

Exit codes: 0 ok, 2 validation errors (nothing written).

Function index:
read_workbook, build_plan, norm_url, diff_plan, snapshot_to_workbook, main
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

LEVEL_HEADERS = ["Top Level", "2nd Level", "3rd Level"]
OPTIONAL_HEADERS = ["Url", "OpenInNewTab", "Enabled"]
FALSE_WORDS = {"n", "no", "false", "0"}
TRUE_WORDS = {"y", "yes", "true", "1"}
OP_KINDS = ["create", "update", "unchanged", "order-drift", "prune"]


def read_workbook(path: str | Path) -> list[dict]:
    """Read the first sheet into row dicts keyed by canonical header; raises ValueError on missing level headers."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Workbook is empty.")
    lookup = {str(h).strip().casefold(): i for i, h in enumerate(rows[0]) if h is not None}
    missing = [h for h in LEVEL_HEADERS if h.casefold() not in lookup]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}. Expected {LEVEL_HEADERS}.")
    out = []
    for number, raw in enumerate(rows[1:], start=2):
        row = {"row": number}
        for header in LEVEL_HEADERS + OPTIONAL_HEADERS:
            idx = lookup.get(header.casefold())
            value = raw[idx] if idx is not None and idx < len(raw) else None
            row[header] = "" if value is None else str(value).strip()
        if any(row[h] for h in LEVEL_HEADERS):
            out.append(row)
    return out


def _check_url(url: str) -> str | None:
    """Return an error fragment for an unusable URL, or None when acceptable."""
    if "#" in url:
        return "contains a '#' fragment"
    if url.startswith("/") and not url.startswith("//"):
        return None
    parsed = urlparse(url)
    if parsed.scheme in ("http", "https") and parsed.netloc:
        return None
    return "must be site-relative (/...) or http(s)"


def _flag(value: str, words: set[str]) -> bool:
    """True when value is one of the accepted words (case-insensitive)."""
    return value.strip().casefold() in words


def build_plan(rows: list[dict], allow_missing_urls: bool = False) -> dict:
    """Validate rows and build the nested plan: {"nodes", "errors", "warnings"}."""
    errors: list[str] = []
    warnings: list[str] = []
    roots: list[dict] = []
    index: dict[tuple, dict] = {}
    disabled: list[tuple] = []

    for row in rows:
        levels = [row[h] for h in LEVEL_HEADERS]
        depth = max(i for i, v in enumerate(levels) if v) + 1
        path = tuple(levels[:depth])
        label = f"row {row['row']} ({' > '.join(path)})"
        if any(not v for v in path):
            errors.append(f"{label}: hierarchy gap - a level is blank above a populated level.")
            continue
        if any("\ufffd" in v for v in path):
            errors.append(f"{label}: encoding problem - title contains a U+FFFD replacement character.")
            continue
        if _flag(row["Enabled"], FALSE_WORDS):
            disabled.append(path)
            continue
        if row["Url"]:
            problem = _check_url(row["Url"])
            if problem:
                errors.append(f"{label}: url {problem}: {row['Url']}")
                continue
        existing = index.get(path)
        if existing is not None and not existing["implicit"]:
            errors.append(f"{label}: duplicate full path.")
            continue
        node = existing or {"title": path[-1], "url": "", "openInNewTab": False, "children": [], "implicit": False}
        node.update(url=row["Url"], openInNewTab=_flag(row["OpenInNewTab"], TRUE_WORDS), implicit=False)
        if existing is None:
            parent = _ensure_parents(path, index, roots, warnings)
            (parent["children"] if parent else roots).append(node)
            index[path] = node

    dead = {p for p in index if any(p[: len(d)] == d for d in disabled)}
    if dead:
        _drop(roots, dead)
    _check_leaves(roots, [], errors, warnings, allow_missing_urls)
    _strip_internal(roots)
    return {"nodes": roots, "errors": errors, "warnings": warnings}


def _ensure_parents(path: tuple, index: dict, roots: list, warnings: list) -> dict | None:
    """Return the parent node for path, creating implicit ancestors (with a warning) as needed."""
    parent = None
    for depth in range(1, len(path)):
        key = path[:depth]
        if key not in index:
            node = {"title": key[-1], "url": "", "openInNewTab": False, "children": [], "implicit": True}
            (parent["children"] if parent else roots).append(node)
            index[key] = node
            warnings.append(f"implicit parent created: {' > '.join(key)}")
        parent = index[key]
    return parent


def _drop(nodes: list, dead: set, trail: tuple = ()) -> None:
    """Remove nodes whose full path is in dead (children go with them)."""
    for node in list(nodes):
        path = trail + (node["title"],)
        if path in dead:
            nodes.remove(node)
        else:
            _drop(node["children"], dead, path)


def _check_leaves(nodes: list, trail: list, errors: list, warnings: list, allow: bool) -> None:
    """Leaves need a URL; group nodes may be linkless headers."""
    for node in nodes:
        path = trail + [node["title"]]
        if node["children"]:
            _check_leaves(node["children"], path, errors, warnings, allow)
        elif not node["url"]:
            (warnings if allow else errors).append(f"{' > '.join(path)}: leaf has no url.")


def _strip_internal(nodes: list) -> None:
    """Drop bookkeeping keys before the plan is serialised."""
    for node in nodes:
        node.pop("implicit", None)
        _strip_internal(node["children"])


def norm_url(url: str | None, site_url: str = "") -> str:
    """Normalise a URL for comparison: linkless->"", strip same-site origin, casefold, trim trailing '/'."""
    if not url or "linkless.header" in url.casefold():
        return ""
    text = url.strip()
    if site_url:
        site = urlparse(site_url)
        origin = f"{site.scheme}://{site.netloc}".casefold()
        if text.casefold().startswith(origin):
            text = text[len(origin):] or "/"
    return text.casefold().rstrip("/") or "/"


def diff_plan(plan: dict, snapshot: dict, site_url: str = "") -> dict:
    """Compare plan nodes with a live snapshot and return {"ops": [...], "summary": {...}} in preorder."""
    live_nodes = snapshot.get("topNav") or snapshot.get("quickLaunch") or []
    ops: list[dict] = []
    _walk(plan["nodes"], live_nodes, [], ops, site_url)
    summary = {kind: sum(1 for o in ops if o["op"] == kind) for kind in OP_KINDS}
    return {"ops": ops, "summary": summary}


def _walk(plan_nodes: list, live_nodes: list, parent_path: list, ops: list, site_url: str) -> None:
    """Emit ops for one sibling group, then recurse into each planned node's children."""
    live_by_title = {n["title"]: (i, n) for i, n in enumerate(live_nodes)}
    previous: list | None = None
    last_index = -1
    for node in plan_nodes:
        path = parent_path + [node["title"]]
        match = live_by_title.get(node["title"])
        if match:
            position, live = match
            same = norm_url(live.get("url"), site_url) == norm_url(node["url"], site_url)
            op = {"op": "unchanged" if same else "update", "path": path, "id": live["id"]}
            if not same:
                op["url"] = node["url"]
            ops.append(op)
            if position < last_index:
                ops.append({"op": "order-drift", "path": path, "id": live["id"]})
            last_index = max(last_index, position)
            children_live = live.get("children", [])
        else:
            ops.append({"op": "create", "path": path, "parentPath": parent_path or None, "previousPath": previous,
                        "url": node["url"], "openInNewTab": node["openInNewTab"]})
            children_live = []
        _walk(node["children"], children_live, path, ops, site_url)
        previous = path
    planned = {n["title"] for n in plan_nodes}
    for live in live_nodes:
        if live["title"] not in planned:
            ops.append({"op": "prune", "path": parent_path + [live["title"]], "id": live["id"],
                        "title": live["title"]})


def snapshot_to_workbook(snapshot: dict, out: str | Path) -> None:
    """Write a snapshot as xlsx or csv using the workbook columns (one row per node, parents included)."""
    header = LEVEL_HEADERS + OPTIONAL_HEADERS
    rows: list[list] = []

    def emit(nodes: list, trail: list) -> None:
        for node in nodes:
            path = trail + [node["title"]]
            url = node.get("url") or ""
            url = "" if "linkless.header" in url.casefold() else url
            rows.append((path + [""] * 3)[:3] + [url, "", ""])
            emit(node.get("children", []), path)

    emit(snapshot.get("topNav") or snapshot.get("quickLaunch") or [], [])
    out = Path(out)
    if out.suffix.lower() == ".csv":
        with out.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(header)
            writer.writerows(rows)
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Navigation Snapshot"
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(out)


def main(argv: list[str] | None = None) -> int:
    """CLI entry point."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="cmd", required=True)
    p_plan = sub.add_parser("plan")
    p_plan.add_argument("--workbook", required=True)
    p_plan.add_argument("--out", required=True)
    p_plan.add_argument("--allow-missing-urls", action="store_true")
    p_diff = sub.add_parser("diff")
    p_diff.add_argument("--plan", required=True)
    p_diff.add_argument("--snapshot", required=True)
    p_diff.add_argument("--out", required=True)
    p_diff.add_argument("--site-url", default="")
    p_wb = sub.add_parser("to-workbook")
    p_wb.add_argument("--snapshot", required=True)
    p_wb.add_argument("--out", required=True)
    args = parser.parse_args(argv)

    if args.cmd == "plan":
        try:
            plan = build_plan(read_workbook(args.workbook), args.allow_missing_urls)
        except ValueError as exc:
            print(f"ERROR: {exc}", file=sys.stderr)
            return 2
        for warning in plan["warnings"]:
            print(f"WARNING: {warning}", file=sys.stderr)
        if plan["errors"]:
            for error in plan["errors"]:
                print(f"ERROR: {error}", file=sys.stderr)
            return 2
        Path(args.out).write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps({"nodes": len(plan["nodes"]), "warnings": len(plan["warnings"])}))
    elif args.cmd == "diff":
        plan = json.loads(Path(args.plan).read_text(encoding="utf-8"))
        snapshot = json.loads(Path(args.snapshot).read_text(encoding="utf-8-sig"))
        site_url = args.site_url or snapshot.get("siteUrl", "")
        result = diff_plan(plan, snapshot, site_url)
        result.update(siteUrl=site_url, location="quickLaunch" if "quickLaunch" in snapshot else "topNav")
        Path(args.out).write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(result["summary"]))
    else:
        snapshot = json.loads(Path(args.snapshot).read_text(encoding="utf-8-sig"))
        snapshot_to_workbook(snapshot, args.out)
        print(json.dumps({"written": args.out}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
